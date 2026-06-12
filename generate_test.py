import docx
from docx.shared import Pt
import os

def create_test_docx():
    doc = docx.Document()
    
    # Title
    doc.add_heading('안녕하세요 Test Doc', 0)
    
    # Paragraph
    p = doc.add_paragraph('이 문서는 파이썬으로 자동 생성되었습니다. ')
    p.add_run('굵은 글씨(Bold)').bold = True
    p.add_run('와 ')
    p.add_run('이탤릭(Italic)').italic = True
    p.add_run('입니다.')
    
    # Heading 1
    doc.add_heading('1. 목록 테스트', level=1)
    doc.add_paragraph('첫 번째 항목', style='List Bullet')
    doc.add_paragraph('두 번째 항목', style='List Bullet')
    
    # Table headers
    doc.add_heading('2. 표 테스트 (Table)', level=1)
    table = doc.add_table(rows=3, cols=2)
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '이름'
    hdr_cells[1].text = '직업'
    
    row1 = table.rows[1].cells
    row1[0].text = '홍길동'
    row1[1].text = '의적'
    
    row2 = table.rows[2].cells
    row2[0].text = '임꺽정'
    row2[1].text = '도적'
    
    path = os.path.abspath("test_korean.docx")
    doc.save(path)
    print(f"Created: {path}")
    return path

if __name__ == "__main__":
    create_test_docx()
