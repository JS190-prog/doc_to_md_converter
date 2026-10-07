import os
from itertools import zip_longest

from src.utils.md_table import rows_to_markdown


def _merge_cached_and_formula_rows(value_rows, formula_rows):
    """Prefer cached values; fall back to the formula text for cells whose
    cache is empty (workbooks never recalculated/saved by Excel)."""
    for values, formulas in zip_longest(value_rows, formula_rows, fillvalue=()):
        merged = []
        for value, formula in zip_longest(values, formulas, fillvalue=None):
            if value is None and isinstance(formula, str) and formula.startswith("="):
                merged.append(formula)
            else:
                merged.append(value)
        yield merged


def convert_xlsx_to_md(file_path):
    """
    Converts a .xlsx file to Markdown.
    Each sheet is emitted as a section with a markdown table.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ImportError("openpyxl is required for .xlsx conversion") from e

    values_wb = load_workbook(file_path, data_only=True, read_only=True)
    formulas_wb = load_workbook(file_path, data_only=False, read_only=True)
    output = []
    try:
        for sheet in values_wb.worksheets:
            output.append(f"## {sheet.title}\n\n")
            formula_sheet = formulas_wb[sheet.title]
            rows = list(
                _merge_cached_and_formula_rows(
                    sheet.iter_rows(values_only=True),
                    formula_sheet.iter_rows(values_only=True),
                )
            )
            output.append(rows_to_markdown(rows) if rows else "_(empty sheet)_\n\n")
    finally:
        values_wb.close()
        formulas_wb.close()

    return "".join(output).strip() + "\n"
